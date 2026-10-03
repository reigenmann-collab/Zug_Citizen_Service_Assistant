"""Erzeugt das Pruefblatt fuer die Gold-Labels: eval/label_sheet.xlsx (plus .csv als Rueckfallebene).

Die Labels in testset.json stammen von Claude und sind damit nicht unabhaengig von dem System, das
sie pruefen sollen (siehe eval/README.md). Dieses Blatt dient dazu, sie von Hand zu bestaetigen oder
zu korrigieren. Zurueckgelesen wird mit eval/label_sheet_einlesen.py.

Reihenfolge nach Prioritaet: zuerst die Faelle, bei denen das Urteil wirklich etwas entscheidet.
"""
import csv, json
from pathlib import Path

P = Path(__file__).parent

# Prio 1: Ermessen oder strittig - hier entscheidet das Urteil ueber die Bewertung der Pipeline.
PRIO1 = {"D02", "D16", "D17", "D18", "D19", "D23"}
# Prio 2: Zwei Quellen noetig - der Kern des Coverage-Checks.
PRIO2_KAT = {"B_zwei_quellen", "B_zwei_quellen_en", "F_mehrdeutig", "C_hard_routing", "E_falschannahme"}


def prio(t):
    if t["id"] in PRIO1:
        return 1
    if t["kategorie"] in PRIO2_KAT:
        return 2
    return 3


SPALTEN = [
    ("Prio", 6), ("id", 7), ("split", 7), ("kategorie", 20), ("sprache", 9),
    ("frage", 44), ("erwartetes_verhalten", 24), ("gold_kurzantwort", 60),
    ("erwartete_quellen", 26), ("muss_enthalten", 16), ("begruendung", 52),
    ("Gold_OK", 11), ("Korrektur_Gold_Antwort", 44), ("Korrektur_Verhalten", 22), ("Bemerkung", 36),
]
HINWEIS = [
    "Pruefblatt Gold-Labels – Stadt Zug RAG",
    "",
    "Was zu tun ist: Spalte 'Gold_OK' ausfuellen (ja / nein / unsicher).",
    "  ja       – Frage, erwartetes Verhalten und Gold-Antwort sind so richtig.",
    "  nein     – bitte 'Korrektur_Gold_Antwort' und/oder 'Korrektur_Verhalten' ausfuellen.",
    "  unsicher – wird im Bericht getrennt ausgewiesen und nicht als bestaetigt gezaehlt.",
    "",
    "Reihenfolge ist nach Prioritaet sortiert:",
    "  Prio 1 (6 Zeilen)  – Ermessensfaelle und ein strittiges Label. Diese entscheiden, wie die",
    "                       Pipeline bewertet wird. Fuer den Start von Phase 5 genuegen diese.",
    "  Prio 2 (14 Zeilen) – Faelle mit zwei noetigen Quellen, Mehrdeutigkeit, Eskalation,",
    "                       Falschannahme. Kern der Governance-Anforderungen.",
    "  Prio 3 (24 Zeilen) – einfache Faelle; vor Phase 7 pruefen, nicht jetzt.",
    "",
    "Hintergrund: Die Labels hat Claude aus dem indexierten Korpus abgeleitet. Sie sind damit nicht",
    "unabhaengig vom geprueften System: Luecken des Korpus bleiben unentdeckt, weil die Fragen so",
    "formuliert sind, dass ihre Antwort im Korpus steht. Deshalb diese Durchsicht.",
    "",
    "Zurueckgelesen wird mit:  python eval/label_sheet_einlesen.py",
    "Die Datei darf dabei geschlossen sein; nur die Spalten Gold_OK, Korrektur_* und Bemerkung",
    "werden ausgewertet.",
]


def main():
    rows = sorted(json.loads((P / "testset.json").read_text(encoding="utf8")),
                  key=lambda t: (prio(t), t["id"]))

    # CSV als Rueckfallebene (Semikolon, UTF-8 mit BOM fuer Excel)
    with (P / "label_sheet.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([n for n, _ in SPALTEN])
        for t in rows:
            w.writerow([prio(t), t["id"], t["split"], t["kategorie"], t["sprache"], t["frage"],
                        t["erwartetes_verhalten"], t["gold_kurzantwort"], " | ".join(t["erwartete_quellen"]),
                        " | ".join(t["muss_enthalten"]), t["begruendung"], "", "", "", ""])

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl fehlt (pip install openpyxl) – nur label_sheet.csv erzeugt.")
        return

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Anleitung"
    for i, zeile in enumerate(HINWEIS, 1):
        c = ws0.cell(row=i, column=1, value=zeile)
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 105

    ws = wb.create_sheet("Labels")
    kopf = Font(bold=True, color="FFFFFF")
    fill_kopf = PatternFill("solid", fgColor="435075")          # primary-650 wie im UI-Entwurf
    fill_pruef = PatternFill("solid", fgColor="FFF4E8")          # auszufuellende Spalten
    fill_p1 = PatternFill("solid", fgColor="E07031")             # Akzentfarbe fuer Prio 1
    fill_p2 = PatternFill("solid", fgColor="F5F9FF")
    rand = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

    for j, (name, breite) in enumerate(SPALTEN, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = kopf, fill_kopf
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = breite
    ws.row_dimensions[1].height = 28

    pruefspalten = {"Gold_OK", "Korrektur_Gold_Antwort", "Korrektur_Verhalten", "Bemerkung"}
    for i, t in enumerate(rows, 2):
        werte = [prio(t), t["id"], t["split"], t["kategorie"], t["sprache"], t["frage"],
                 t["erwartetes_verhalten"], t["gold_kurzantwort"], " | ".join(t["erwartete_quellen"]),
                 " | ".join(t["muss_enthalten"]), t["begruendung"], "", "", "", ""]
        for j, ((name, _), v) in enumerate(zip(SPALTEN, werte), 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = rand
            if name in pruefspalten:
                c.fill = fill_pruef
            elif name == "Prio":
                c.fill = fill_p1 if prio(t) == 1 else (fill_p2 if prio(t) == 2 else PatternFill())
                c.alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[i].height = 112 if prio(t) == 1 else 92

    dv = DataValidation(type="list", formula1='"ja,nein,unsicher"', allow_blank=True,
                        prompt="ja / nein / unsicher", promptTitle="Gold-Label bestaetigt?")
    ws.add_data_validation(dv)
    sp = get_column_letter(1 + [n for n, _ in SPALTEN].index("Gold_OK"))
    dv.add(f"{sp}2:{sp}{len(rows) + 1}")

    ws.freeze_panes = "F2"          # id/kategorie/sprache bleiben sichtbar
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{len(rows) + 1}"
    wb.save(P / "label_sheet.xlsx")

    n1 = sum(1 for t in rows if prio(t) == 1)
    n2 = sum(1 for t in rows if prio(t) == 2)
    print(f"eval/label_sheet.xlsx und .csv erzeugt: {len(rows)} Zeilen "
          f"(Prio 1: {n1}, Prio 2: {n2}, Prio 3: {len(rows) - n1 - n2})")


if __name__ == "__main__":
    main()
