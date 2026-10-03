"""Liest das ausgefuellte Pruefblatt zurueck in eval/testset.json.

Gelesen wird label_sheet.xlsx, falls vorhanden, sonst label_sheet.csv. Uebernommen werden nur die
Spalten Gold_OK, Korrektur_Gold_Antwort, Korrektur_Verhalten und Bemerkung.

Wirkung je Zeile:
  ja        -> label_quelle = "rene_bestaetigt", geprueft_von/geprueft_am gesetzt
  nein      -> Korrekturen werden uebernommen, label_quelle = "rene_korrigiert"
  unsicher  -> label_quelle = "rene_unsicher"; diese Zeilen werden in der Auswertung getrennt
               ausgewiesen und nicht als bestaetigt gezaehlt
  leer      -> unveraendert (bleibt claude_vorschlag)

Vor dem Schreiben wird eine Sicherung testset_vor_pruefung.json angelegt.
Usage: python eval/label_sheet_einlesen.py [--pruefer "René Eigenmann"] [--trocken]
"""
import argparse, csv, datetime as dt, json, shutil, sys
from pathlib import Path

P = Path(__file__).parent
SPALTEN = ("Gold_OK", "Korrektur_Gold_Antwort", "Korrektur_Verhalten", "Bemerkung")


def aus_xlsx(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    ws = wb["Labels"] if "Labels" in wb.sheetnames else wb.active
    kopf = [c.value for c in ws[1]]
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if zeile and zeile[kopf.index("id")]:
            yield {k: v for k, v in zip(kopf, zeile)}


def aus_csv(p):
    with p.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if r.get("id"):
                yield r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pruefer", default="René Eigenmann")
    ap.add_argument("--trocken", action="store_true", help="nur anzeigen, nichts schreiben")
    a = ap.parse_args()

    xlsx, csvp = P / "label_sheet.xlsx", P / "label_sheet.csv"
    if xlsx.exists():
        zeilen = list(aus_xlsx(xlsx)); quelle = xlsx.name
    elif csvp.exists():
        zeilen = list(aus_csv(csvp)); quelle = csvp.name
    else:
        sys.exit("Kein Pruefblatt gefunden. Zuerst: python eval/label_sheet.py")

    pfad = P / "testset.json"
    T = json.loads(pfad.read_text(encoding="utf8"))
    nach_id = {t["id"]: t for t in T}
    heute = dt.date.today().isoformat()
    zaehler = {"bestaetigt": 0, "korrigiert": 0, "unsicher": 0, "offen": 0}
    unbekannt, protokoll = [], []

    for r in zeilen:
        t = nach_id.get(str(r.get("id", "")).strip())
        if not t:
            unbekannt.append(r.get("id")); continue
        ok = (str(r.get("Gold_OK") or "").strip().lower())
        bem = (str(r.get("Bemerkung") or "").strip() or None)
        if bem:
            t["bemerkung_pruefer"] = bem
        if ok in ("", "none"):
            zaehler["offen"] += 1; continue
        t["geprueft_von"], t["geprueft_am"] = a.pruefer, heute
        if ok.startswith("ja"):
            t["label_quelle"] = "rene_bestaetigt"; zaehler["bestaetigt"] += 1
        elif ok.startswith("unsicher"):
            t["label_quelle"] = "rene_unsicher"; zaehler["unsicher"] += 1
        else:
            neu_a = (str(r.get("Korrektur_Gold_Antwort") or "").strip() or None)
            neu_v = (str(r.get("Korrektur_Verhalten") or "").strip() or None)
            if neu_a:
                t.setdefault("gold_kurzantwort_claude", t["gold_kurzantwort"]); t["gold_kurzantwort"] = neu_a
            if neu_v:
                t.setdefault("erwartetes_verhalten_claude", t["erwartetes_verhalten"]); t["erwartetes_verhalten"] = neu_v
            t["label_quelle"] = "rene_korrigiert"; zaehler["korrigiert"] += 1
            protokoll.append(f"  {t['id']}: Verhalten={neu_v or 'unveraendert'}; Antwort={'neu' if neu_a else 'unveraendert'}")

    print(f"Gelesen aus {quelle}: {len(zeilen)} Zeilen")
    print(f"  bestaetigt: {zaehler['bestaetigt']}, korrigiert: {zaehler['korrigiert']}, "
          f"unsicher: {zaehler['unsicher']}, noch offen: {zaehler['offen']}")
    if protokoll:
        print("Korrekturen:"); print("\n".join(protokoll))
    if unbekannt:
        print(f"  unbekannte IDs uebersprungen: {unbekannt}")
    if a.trocken:
        print("Trockenlauf – testset.json unveraendert."); return
    if zaehler["bestaetigt"] + zaehler["korrigiert"] + zaehler["unsicher"] == 0:
        print("Nichts ausgefuellt – testset.json unveraendert."); return
    sich = P / "testset_vor_pruefung.json"
    if not sich.exists():
        shutil.copy(pfad, sich); print(f"Sicherung: {sich.name}")
    pfad.write_text(json.dumps(T, ensure_ascii=False, indent=1), encoding="utf8")
    print("testset.json aktualisiert.")


if __name__ == "__main__":
    main()
